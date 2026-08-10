# -*- coding: utf-8 -*-
# =============================================================================
# WRSK.MD_IFRS_IRB_RCR_CST_APPL_INF — Veri Validasyonu Excel Raporu
# =============================================================================
# Dataiku Python notebook'unda çalışır:
#   1. 12 Oracle sorgusunu SQLExecutor2 ile sırayla çalıştırır.
#   2. Her sorgunun sonucunu tek bir .xlsx dosyasında AYRI bir sayfaya yazar
#      (1-sonuç ... 12-sonuç) ve başa bir "Özet" sayfası ekler.
#   3. Kod bittiğinde raporu tarayıcıya OTOMATİK indirir.
#
# Kullanım:
#   - ORACLE_CONNECTION değişkenine Dataiku'daki Oracle bağlantı adınızı yazın.
#   - SORGULAR listesindeki örnek sorguların yerine kendi 12 sorgunuzu
#     yapıştırın (sayfa adları ve açıklamalar hazır).
#   - Tüm hücreleri çalıştırın (Cell > Run All). İndirme otomatik başlar;
#     başlamazsa çıktıda görünen linke tıklayın.
#
# Notlar:
#   - Sorguların sonuna ';' KOYMAYIN (JDBC üzerinden ORA-00911 hatası verir).
#   - Bir sorgu hata alırsa rapor durmaz: ilgili sayfaya hata mesajı yazılır,
#     Özet sayfasında DURUM = HATA olarak görünür.
#   - openpyxl gerektirir (Dataiku builtin env'de genelde vardır; yoksa code
#     env'inize ekleyin).
# =============================================================================

# %% 1) Ayarlar ve 12 sorgu

import base64
import os
import tempfile
import time
import traceback
from datetime import datetime

import pandas as pd

import dataiku
from dataiku import SQLExecutor2
from IPython.display import HTML, display

# Dataiku'daki Oracle bağlantısının adı (Administration > Connections)
ORACLE_CONNECTION = "ORACLE_BAGLANTI_ADI"  # TODO: kendi bağlantı adınızı yazın
# Alternatif: bağlantı adı yerine o bağlantıyı kullanan bir dataset ile de
# açabilirsiniz:  executor = SQLExecutor2(dataset="DATASET_ADI")

RAPOR_ON_EKI = "WRSK_MD_IFRS_IRB_RCR_CST_APPL_INF_validasyon"

# (sayfa adı, kontrol noktası açıklaması, SQL)
# Her sorgunun yerine kendi SQL'inizi üç tırnak arasına yapıştırın.
SORGULAR = [
    ("1-sonuç",
     "MUST_NO (müşteri numarası) uzunluk / NULL / '0' / geçersiz karakter kontrolleri",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 1. sorgunuzu buraya yapıştırın
     """),

    ("2-sonuç",
     "TAR alanına göre aylık kayıt adedi (ADET_TAR) ve önceki aya göre % değişim",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 2. sorgunuzu buraya yapıştırın
     """),

    ("3-sonuç",
     "BNK_KOD, KYNK_KOD, KANAL_KOD, BSVR_TIP_KOD, KRD_TIP_KOD, ONAY_DRM_KOD, "
     "IKMT_DRM_KOD, IS_SZLSM_TIP_KOD kontrolleri",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 3. sorgunuzu buraya yapıştırın
     """),

    ("4-sonuç",
     "TAR, BSVR_TAR, AKTAR_TAR tarih alanları NULL / format kontrolleri",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 4. sorgunuzu buraya yapıştırın
     """),

    ("5-sonuç",
     "IS_AKIS_NO, IS_SURUM_NO uzunluk / NULL / geçersiz karakter kontrolleri",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 5. sorgunuzu buraya yapıştırın
     """),

    ("6-sonuç",
     "BSVR_TUT negatif / NULL kontrolleri",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 6. sorgunuzu buraya yapıştırın
     """),

    ("7-sonuç",
     "MUST_NO + TAR kombinasyonunda tekrar eden kayıt kontrolü",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 7. sorgunuzu buraya yapıştırın
     """),

    ("8-sonuç",
     "Alan bazlı toplam kayıt sayısı ve trimlenmemiş değer kontrolü",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 8. sorgunuzu buraya yapıştırın
     """),

    ("9-sonuç",
     "NET_GLR_TUT negatif / NULL kontrolleri",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 9. sorgunuzu buraya yapıştırın
     """),

    ("10-sonuç",
     "IKMT_SURE, IS_CLSM_SURE, ESKI_IS_CLSM_SURE kontrolleri",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 10. sorgunuzu buraya yapıştırın
     """),

    ("11-sonuç",
     "REF_NO ve BSVR_SKOR NULL kontrolleri",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 11. sorgunuzu buraya yapıştırın
     """),

    ("12-sonuç",
     "KRD_VADE uzunluk / NULL / geçersiz karakter kontrolleri",
     """
     SELECT 1 AS ORNEK_KOLON FROM DUAL  -- TODO: 12. sorgunuzu buraya yapıştırın
     """),
]

# %% 2) Sorguları çalıştır

executor = SQLExecutor2(connection=ORACLE_CONNECTION)

sonuclar = []  # her öğe: dict(sayfa, aciklama, df, hata, satir, sure_sn)
for sira, (sayfa, aciklama, sql) in enumerate(SORGULAR, start=1):
    print(f"[{sira}/{len(SORGULAR)}] {sayfa} çalışıyor...", end=" ", flush=True)
    baslangic = time.time()
    try:
        df = executor.query_to_df(sql)
        hata = None
        print(f"OK ({len(df)} satır, {time.time() - baslangic:.1f} sn)")
    except Exception:
        df = None
        hata = traceback.format_exc()
        print("HATA! (rapor devam ediyor, detay ilgili sayfada)")
    sonuclar.append({
        "sayfa": sayfa,
        "aciklama": aciklama,
        "df": df,
        "hata": hata,
        "satir": 0 if df is None else len(df),
        "sure_sn": round(time.time() - baslangic, 1),
    })

# %% 3) Excel raporunu oluştur (her sorgu ayrı sayfa + Özet)

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASLIK_FONT = Font(bold=True, color="FFFFFF")
BASLIK_DOLGU = PatternFill("solid", fgColor="305496")


def sayfayi_bicimlendir(ws, df):
    """Başlık satırını renklendirir, filtre ekler, kolon genişliğini ayarlar."""
    ws.freeze_panes = "A2"
    if len(df) > 0:
        ws.auto_filter.ref = ws.dimensions
    for j, kolon in enumerate(df.columns, start=1):
        hucre = ws.cell(row=1, column=j)
        hucre.font = BASLIK_FONT
        hucre.fill = BASLIK_DOLGU
        hucre.alignment = Alignment(vertical="center")
        ornekler = df[kolon].head(200).astype(str).map(len)
        genislik = max(int(ornekler.max()) if len(ornekler) else 0, len(str(kolon)))
        ws.column_dimensions[get_column_letter(j)].width = min(genislik + 3, 60)


zaman_damgasi = datetime.now().strftime("%Y%m%d_%H%M%S")
dosya_adi = f"{RAPOR_ON_EKI}_{zaman_damgasi}.xlsx"
dosya_yolu = os.path.join(tempfile.gettempdir(), dosya_adi)

ozet = pd.DataFrame([
    {
        "SORGU NO": i + 1,
        "SAYFA": s["sayfa"],
        "KONTROL NOKTASI": s["aciklama"],
        "DURUM": "HATA" if s["hata"] else "OK",
        "SATIR SAYISI": s["satir"],
        "SÜRE (SN)": s["sure_sn"],
    }
    for i, s in enumerate(sonuclar)
])

with pd.ExcelWriter(dosya_yolu, engine="openpyxl") as writer:
    ozet.to_excel(writer, sheet_name="Özet", index=False)
    sayfayi_bicimlendir(writer.sheets["Özet"], ozet)

    for s in sonuclar:
        if s["hata"] is None:
            df = s["df"]
        else:  # sorgu hata aldıysa hata metnini sayfaya yaz
            df = pd.DataFrame({"HATA": s["hata"].splitlines()})
        sayfa_adi = s["sayfa"][:31]  # Excel sayfa adı üst sınırı 31 karakter
        df.to_excel(writer, sheet_name=sayfa_adi, index=False)
        sayfayi_bicimlendir(writer.sheets[sayfa_adi], df)

print(f"Rapor oluşturuldu: {dosya_yolu} ({os.path.getsize(dosya_yolu) / 1024:.0f} KB)")
display(ozet)

# %% 4) Raporu otomatik indir


def raporu_indir(yol):
    """Dosyayı base64'e çevirip indirme linki olarak gömer ve otomatik tıklar.

    Klasik Jupyter arayüzünde (Dataiku notebook'ları) indirme otomatik başlar.
    JupyterLab arayüzü <script> etiketini engellerse çıktıdaki linke elle
    tıklamanız yeterli.
    """
    with open(yol, "rb") as f:
        b64 = base64.b64encode(f.read()).decode()
    ad = os.path.basename(yol)
    link_id = f"rapor_indir_{zaman_damgasi}"
    display(HTML(
        f'<a id="{link_id}" download="{ad}" '
        f'href="data:application/vnd.openxmlformats-officedocument.'
        f'spreadsheetml.sheet;base64,{b64}" '
        f'style="font-size:14px;font-weight:bold">'
        f'&#128229; {ad} — otomatik inmediyse buraya tıklayın</a>'
        f'<script>document.getElementById("{link_id}").click();</script>'
    ))


raporu_indir(dosya_yolu)

# %% 5) (Opsiyonel) Çok büyük raporlar için Managed Folder alternatifi

# Rapor dosyası çok büyürse (ör. >50 MB) base64 ile notebook çıktısına gömmek
# yerine bir Dataiku Managed Folder'a yazıp klasör arayüzünden indirmek daha
# sağlıklıdır. Bunun için Flow'da bir folder oluşturup aşağıyı açın:
#
# klasor = dataiku.Folder("RAPOR_KLASORU")  # folder adı veya id'si
# with open(dosya_yolu, "rb") as f:
#     klasor.upload_stream(dosya_adi, f)
# print("Managed folder'a yüklendi:", dosya_adi)
